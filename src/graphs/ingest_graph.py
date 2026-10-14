"""G0 IngestionGraph: load -> split -> tag -> embed_upsert -> graph edges."""
from __future__ import annotations
import os, hashlib, base64
from pathlib import Path
from typing import TypedDict, Literal
from langgraph.graph import StateGraph, END
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from src.config.settings import get_settings


class IngestState(TypedDict):
    file: str
    kind: str                      # pdf | slides | video_txt
    raw_docs: list
    units: list                    # list[dict citation payload + text]
    errors: list


def _kind(path: str) -> str:
    p = path.lower()
    if p.endswith((".pptx", ".ppt")):
        return "slides"
    if p.endswith((".txt", ".srt", ".vtt")):
        return "video_txt"
    if p.endswith((".jpg", ".jpeg", ".png")):
        return "image"
    return "pdf"


def _clean_meta(md: dict) -> dict:
    """Keep only Chroma-safe scalar metadata; drop nested Docling blobs."""
    return {k: v for k, v in md.items()
            if v is None or isinstance(v, (str, int, float, bool))}


def _describe_with_vlm(path: str, filename: str) -> str:
    """Ask a vision model what the image/diagram shows. Returns '' on any failure
    so the caller can fall back to a filename-only stub."""
    try:
        from src.config.settings import get_settings
        from src.chains.models import get_vision
        from langchain_core.messages import HumanMessage
        s = get_settings()
        if not s.ENABLE_VLM:
            return ""
        data_url = "data:image/%s;base64,%s" % (
            (path.rsplit(".", 1)[-1] or "png").lower(),
            base64.b64encode(Path(path).read_bytes()).decode())
        msg = HumanMessage(content=[
            {"type": "text", "text": (
                "Describe this image or diagram for a study index. In 2-4 sentences: what it "
                "depicts, the labelled parts/components, and any relationship or process shown. "
                "Transcribe short text labels. If it is purely decorative, say so.")},
            {"type": "image_url", "image_url": {"url": data_url}},
        ])
        from src.chains.prompts import _content
        return _content(get_vision().invoke([msg])).strip()
    except Exception:
        return ""


def _image_stub(path: str) -> str:
    """Last-resort description so the file is still searchable/quotable."""
    return f"Image: {os.path.basename(path)} (no description available)."


def _load_images(path: str) -> list:
    """Docling OCR/layout first, VLM description for the picture itself, stub last."""
    from langchain_core.documents import Document
    filename = os.path.basename(path)
    docs: list[Document] = []
    try:
        from langchain_docling import DoclingLoader
        for d in DoclingLoader(file_path=path).load():
            d.metadata = _clean_meta(d.metadata)
            d.metadata.setdefault("source_file", filename)
            d.metadata.setdefault("content_type", "image")
            if d.page_content.strip():
                docs.append(d)
    except Exception:
        pass  # Docling/clipboard/OCR failures are non-fatal — VLM path below still runs

    desc = _describe_with_vlm(path, filename)
    if desc:
        docs.append(Document(page_content=f"{filename}: {desc}",
                             metadata={"source_file": filename,
                                       "content_type": "figure" if len(docs) == 0 else "image_vlm"}))
    if not docs:
        docs.append(Document(page_content=_image_stub(path),
                             metadata={"source_file": filename, "content_type": "image"}))
    return docs


def n_load(state: IngestState) -> IngestState:
    path = state["file"]
    kind = _kind(path)
    state["kind"] = kind
    docs: list[Document] = []
    try:
        if kind == "slides":
            from pptx import Presentation
            prs = Presentation(path)
            for i, s in enumerate(prs.slides, 1):
                text = "\n".join(sh.text for sh in s.shapes if hasattr(sh, "text"))
                docs.append(Document(page_content=f"Slide {i}: {text}",
                                     metadata={"source_file": os.path.basename(path),
                                               "slide": i, "content_type": "slide"}))
        elif kind == "video_txt":
            txt = open(path, encoding="utf-8", errors="ignore").read()
            docs.append(Document(page_content=txt,
                                 metadata={"source_file": os.path.basename(path),
                                           "content_type": "video", "t_start": 0.0}))
        elif kind == "image":
            docs = _load_images(path)
        else:  # pdf via Docling (primary)
            try:
                from langchain_docling import DoclingLoader
                loader = DoclingLoader(file_path=path)
                docs = loader.load()
                for d in docs:
                    # Docling attaches a nested `docling_metadata` dict that Chroma rejects
                    d.metadata = _clean_meta(d.metadata)
                    d.metadata.setdefault("source_file", os.path.basename(path))
                    d.metadata.setdefault("content_type", "textbook")
                    if "page" not in d.metadata and "page_number" in d.metadata:
                        d.metadata["page"] = d.metadata.pop("page_number")
            except Exception:
                from pypdf import PdfReader
                r = PdfReader(path)
                for i, pg in enumerate(r.pages, 1):
                    docs.append(Document(page_content=pg.extract_text() or "",
                                         metadata={"source_file": os.path.basename(path),
                                                   "page": i, "content_type": "textbook"}))
    except Exception as e:
        state["errors"].append(str(e))
    state["raw_docs"] = docs
    return state


def n_split(state: IngestState) -> IngestState:
    s = get_settings()
    if state["kind"] == "slides":      # 1 slide = 1 unit, no further split
        state["units"] = [{"text": d.page_content, **d.metadata, "topic": "general"}
                          for d in state["raw_docs"]]
        return state
    sp = RecursiveCharacterTextSplitter(chunk_size=s.CHUNK_SIZE, chunk_overlap=s.CHUNK_OVERLAP)
    chunks = sp.split_documents(state["raw_docs"])
    state["units"] = [{"text": c.page_content, **c.metadata, "topic": c.metadata.get("topic", "general")}
                      for c in chunks]
    return state


def n_tag(state: IngestState) -> IngestState:
    """Heuristic tagger (offline-safe). LLM tag_chain used when API keys exist."""
    s = get_settings()
    use_llm = bool(s.GOOGLE_API_KEY or s.OPENAI_API_KEY)
    if use_llm:
        try:
            from src.chains.prompts import get_tag_chain
            chain = get_tag_chain()
            for u in state["units"][:50]:  # cap LLM cost
                try:
                    cu = chain.invoke({"source_file": u.get("source_file", ""),
                                       "page": u.get("page"), "slide": u.get("slide"),
                                       "t_start": u.get("t_start"), "t_end": u.get("t_end"),
                                       "content": u["text"][:2000]})
                    u.update(cu.model_dump(exclude_none=True))
                except Exception:
                    pass
            return state
        except Exception:
            pass
    for u in state["units"]:  # fallback tags
        u.setdefault("difficulty", "medium")
        u.setdefault("subtopic", "")
        u.setdefault("concept", u.get("topic", "general"))
    return state


def n_upsert(state: IngestState) -> IngestState:
    from src.stores.vector import add_documents
    docs = [Document(page_content=u["text"],
                     metadata=_clean_meta({k: v for k, v in u.items() if k != "text"}))
            for u in state["units"] if u.get("text", "").strip()]
    if docs:
        add_documents(docs)
    return state


def n_graph(state: IngestState) -> IngestState:
    from src.stores.graph import add_topic_edge
    for u in state["units"]:
        if u.get("concept"):
            add_topic_edge(u["concept"], prereq=u.get("subtopic", ""),
                           topic=u.get("topic", "general"))
    return state


def build_ingest_graph():
    g = StateGraph(IngestState)
    g.add_node("load", n_load); g.add_node("split", n_split)
    g.add_node("tag", n_tag); g.add_node("upsert", n_upsert); g.add_node("graph", n_graph)
    g.set_entry_point("load")
    g.add_edge("load", "split"); g.add_edge("split", "tag")
    g.add_edge("tag", "upsert"); g.add_edge("upsert", "graph"); g.add_edge("graph", END)
    return g.compile()


def ingest_file(path: str) -> dict:
    return build_ingest_graph().invoke({"file": path, "kind": "", "raw_docs": [], "units": [], "errors": []})
